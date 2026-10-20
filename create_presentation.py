import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_BG = RGBColor(246, 241, 234)         # Areia clara / sand
    COLOR_CARD = RGBColor(255, 255, 255)       # Branco puro para cards
    COLOR_CARD_BORDER = RGBColor(218, 208, 195) # Borda suave
    COLOR_EARTH = RGBColor(44, 36, 22)         # Terra escura / texto principal
    COLOR_PRIMARY = RGBColor(139, 94, 60)      # Ocre caatinga / primário
    COLOR_SECONDARY = RGBColor(74, 124, 89)    # Verde caatinga / biofísico
    COLOR_ACCENT = RGBColor(190, 135, 30)      # Amarelo/dourado seco legível
    COLOR_MUTED = RGBColor(115, 105, 95)       # Texto secundário
    COLOR_HYDRO = RGBColor(35, 110, 145)       # Azul hídrico (H2V)
    COLOR_DARK_BAR = RGBColor(38, 30, 20)      # Fundo escuro para cabeçalhos ou destaques

    def set_slide_background(slide, color=COLOR_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="PLATAFORMA PITS • PROJETO FUNDECI/2025.0016"):
        # Header category
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY

        # Header main title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_EARTH

        # Accent line under header
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.fill.background()

    def add_footer(slide, current_slide, total_slides=12):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"PITS — Framework HAMILTON • Dados Auditados: IBGE (2022/2025) | MapBiomas Col. 11 (2025) | SFB/IFN | INMET                Slide {current_slide} de {total_slides}"
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_MUTED

    def create_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # ==========================================
    # SLIDE 1: CAPA EXECUTIVA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_BG)

    # Decorative banner container
    c_banner = create_card(s1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=COLOR_DARK_BAR, border_color=COLOR_PRIMARY)
    
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "PROJETO FUNDECI / 2025.0016 • BANCO DO NORDESTE"
    p0.font.name = "Arial"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(18)

    p1 = tf.add_paragraph()
    p1.text = "PITS — Plataforma de Inteligência Territorial Socioenergética"
    p1.font.name = "Arial"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Framework HAMILTON: Integração Comunidade–Biomassa para o Planejamento Territorial Sustentável da Caatinga e Transição Energética Justa"
    p2.font.name = "Arial"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(220, 215, 205)
    p2.space_after = Pt(32)

    p3 = tf.add_paragraph()
    p3.text = "Apresentação Executiva para Financiadores, Parceiros Estratégicos e Tomadores de Decisão\nTerritório Piloto: 8 Municípios do Semiárido Cearense • Modelagem Científica 100% Auditada"
    p3.font.name = "Arial"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(180, 170, 160)

    # ==========================================
    # SLIDE 2: A TESE DE INVESTIMENTO
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "A Tese de Investimento: Por Que a PITS Transforma a Alocação de Recursos?")
    add_footer(s2, 2)

    # Card 1: O Problema
    create_card(s2, Inches(0.8), Inches(1.7), Inches(3.6), Inches(5.0))
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(3.2), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. O Paradoxo do Semiárido"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(180, 50, 40)
    p.space_after = Pt(12)
    
    p = tf.add_paragraph()
    p.text = "• A Caatinga abriga a maior densidade populacional semiárida do planeta (~28 milhões de pessoas).\n• Alta dependência de lenha predatória para energia domiciliar e cerâmica.\n• Políticas públicas tradicionais olham recursos florestais e carências sociais de forma isolada, gerando ineficiência nos investimentos."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # Card 2: A Solução PITS
    create_card(s2, Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.0))
    tb = s2.shapes.add_textbox(Inches(5.0), Inches(1.9), Inches(3.3), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. A Resposta da PITS"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(12)
    
    p = tf.add_paragraph()
    p.text = "• Integração Matemática Pioneira: cruza o Potencial Energético Local (PEL) com o Índice de Vulnerabilidade Social e Climático (IVS*).\n• Gera o IPSE (Índice de Potencial Socioenergético), apontando exatamente onde o investimento gera o maior retorno combinado.\n• Transforma o extrativismo predatório em bioeconomia circular sustentável."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # Card 3: Valor para o Financiador
    create_card(s2, Inches(8.9), Inches(1.7), Inches(3.6), Inches(5.0))
    tb = s2.shapes.add_textbox(Inches(9.1), Inches(1.9), Inches(3.2), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Retorno do Financiamento"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    
    p = tf.add_paragraph()
    p.text = "• Redução do Risco de Crédito/Subvenção: decisões fundamentadas em dados públicos oficiais e auditáveis.\n• Simulação de Cenários 10 Anos: visualização prévia do impacto antes de liberar o desembolso.\n• Conformidade ESG/ODS: alinhado aos ODS 7 (Energia Limpa), 10 (Redução das Desigualdades) e 13 (Ação Climática)."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 3: BASES DE DADOS OFICIAIS E GOVERNANÇA
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Rigor Metodológico: 'Zero Dados Inventados' e Rastreabilidade Total")
    add_footer(s3, 3)

    create_card(s3, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.0))
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.1), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Princípio de Auditoria: Todas as variáveis derivam de fontes primárias do Estado brasileiro"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(14)

    fontes = [
        ("MapBiomas Coleção 11 (2025)", "Mapeamento de 40 anos de uso da terra e área de Caatinga nativa em hectares para cálculo do PEL e conservação."),
        ("IBGE Censo Demográfico 2022", "Microdados oficiais de carência domiciliar (água canalizada, esgotamento, coleta de lixo e energia) para o IVS."),
        ("SFB / IFN-CE & Literatura Científica", "Inventário Florestal Nacional do Ceará para produtividade real de biomassa (t/ha/ano) e poder calorífico (PCI)."),
        ("IBGE PEVS & PPM (Pecuária Municipal)", "Volume extraído de lenha nativa e tamanho dos rebanhos bovino, suíno e caprino para simulação de biogás e H2V."),
        ("INMET BDMEP (Série Histórica)", "Precipitação pluviométrica real captada por estações meteorológicas de referência para cálculo da penalidade de seca."),
        ("IPECE / IQM (Ceará)", "Índice de Qualidade do Meio Ambiente e rateio do ICMS Ecológico para o Índice de Governança Local (IGL).")
    ]

    for nome, desc in fontes:
        p_item = tf.add_paragraph()
        p_item.text = f"• {nome}: {desc}"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = COLOR_EARTH
        p_item.space_after = Pt(8)

    # ==========================================
    # SLIDE 4: O MOTOR CIENTÍFICO (IPSE, PEL, IVS)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "O Motor Científico: O Significado Matemático dos Indicadores")
    add_footer(s4, 4)

    # Box Equação Central
    create_card(s4, Inches(0.8), Inches(1.7), Inches(11.733), Inches(1.4), bg_color=RGBColor(240, 235, 225), border_color=COLOR_PRIMARY)
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.1), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "A EQUAÇÃO MESTRA DO IPSE (Artigo Científico PITS / Seção 3.1):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    p = tf.add_paragraph()
    p.text = "IPSE = PEL / IVS*     |     PEL = Ai · Yi · PCI · η     |     IVS = ∑ wk · Ck"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_EARTH

    # 3 Cards explicativos
    cols = [
        ("PEL (Potencial Energético)", COLOR_SECONDARY, "Eq. 1 — em Gigajoules/ano (GJ/ano)", "Multiplica a área nativa real (MapBiomas) pela produtividade de biomassa (Yi), poder calorífico (PCI) e eficiência térmica (η). Mede o teto de energia renovável manejável."),
        ("IVS* (Vulnerabilidade Climática)", RGBColor(180, 50, 40), "Eq. 4 e 5 — Escala de 0 a 1 (Censo 2022)", "Média ponderada das carências domiciliares das famílias rurais, com piso paramétrico e penalidade pelo déficit hídrico histórico. Quanto MAIOR o IVS, mais urgente é a carência."),
        ("IPSE Normalizado (0 a 1)", COLOR_ACCENT, "Eq. 6 e 7 — Indicador Final de Decisão", "Razão entre o potencial de energia e o índice de vulnerabilidade. Normalizado entre os municípios: Baixo (<0,30), Moderado (0,30–0,59) e Alto (≥0,60). Prioriza projetos de alto impacto.")
    ]

    left_pos = 0.8
    for title, col_color, eq, desc in cols:
        create_card(s4, Inches(left_pos), Inches(3.3), Inches(3.64), Inches(3.4))
        tb = s4.shapes.add_textbox(Inches(left_pos + 0.2), Inches(3.5), Inches(3.24), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(4)
        
        p = tf.add_paragraph()
        p.text = eq
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_MUTED
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_EARTH
        
        left_pos += 4.04

    # ==========================================
    # SLIDE 5: DASHBOARD EXECUTIVO 360°
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Dashboard Executivo 360°: Visão Territorial em Tempo Real")
    add_footer(s5, 5)

    elements = [
        ("1. Velocímetro IPSE Oficial", "Exibe o IPSE normalizado e bruto na tela inicial. Geometria e contrastes calibrados com feedback cromático imediato do status do município."),
        ("2. Quatro Indicadores de Apoio", "Subíndices setoriais: IB (Biomassa nativa), IVS (Vulnerabilidade invertida), IIE (Infraestrutura solar e SAFs) e IGL (Governança e resíduos)."),
        ("3. Mapa Territorial Interativo", "Camadas GIS da Caatinga: limites municipais, corredores ecológicos, zonas de manejo comunitário e alertas de desmatamento."),
        ("4. Série Histórica MapBiomas 40 Anos", "Evolução temporal (1985–2025) dividida entre vegetação nativa preservada, áreas agropecuárias e malha urbana."),
        ("5. Monitoramento SAD-Caatinga", "Contabilização oficial de laudos e hectares desmatados, permitindo intervenções rápidas de fiscalização e proteção florestal."),
        ("6. Alertas e Projeção Decisória", "Motor de regras que emite diagnósticos automáticos (ex.: 'déficit crítico de água') e projeta a curva de 10 anos.")
    ]

    row_h = 1.4
    for i, (title, desc) in enumerate(elements):
        col = i % 2
        row = i // 2
        l = Inches(0.8 + col * 5.95)
        t = Inches(1.7 + row * 1.65)
        create_card(s5, l, t, Inches(5.75), Inches(1.45))
        tb = s5.shapes.add_textbox(l + Inches(0.2), t + Inches(0.15), Inches(5.35), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.space_after = Pt(4)
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 6: MÓDULO I — BIOFÍSICO-ECOLÓGICO
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Módulo I: Biofísico-Ecológico e Manejo Sustentável da Caatinga")
    add_footer(s6, 6)

    create_card(s6, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb1 = s6.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Diagnóstico Biofísico do Território"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(12)
    p = tf1.add_paragraph()
    p.text = "• Cobertura Vegetal Nativa (% e ha): Mede o estoque florestal remanescente da Caatinga pelo MapBiomas 2025.\n• Densidade de Biomassa (t/ha): Estoque em pé mensurado pelo Inventário Florestal Nacional (SFB/IFN-CE).\n• Extração Anual de Lenha (m³/ano): Dados da PEVS/IBGE identificando o nível de pressão sobre a mata nativa.\n• Monitoramento Pluviométrico (INMET): Conexão direta com a estação meteorológica mais próxima para cálculo de resiliência a secas."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    create_card(s6, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb2 = s6.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Catálogo Botânico & PNDBio"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    p = tf2.add_paragraph()
    p.text = "• Catálogo Botânico Interativo: Espécies nativas-chave da Caatinga (Jurema-preta, Marmeleiro, Catingueira, Sabiá) com detalhamento de turnos de corte sustentável e capacidade de rebrota.\n• Zoneamento de Paisagens Multifuncionais: Separa áreas de preservação permanente (APP) e reservas legais de áreas aptas ao manejo florestal comunitário.\n• Alinhamento à PNDBio: Cumpre os eixos estratégicos do Plano Nacional de Desenvolvimento da Bioeconomia."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 7: MÓDULO II — HIDROGÊNIO RENOVÁVEL (H2V)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Módulo II: Simulador de Hidrogênio Renovável (H2) e Bioeconomia Circular")
    add_footer(s7, 7)

    create_card(s7, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Rota A: Gaseificação de Biomassa"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "Reação Termoquímica:\nCHmOn + (2-n)H2O → CO2 + (2-n+m/2)H2"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)
    p = tf1.add_paragraph()
    p.text = "• Transforma resíduos florestais e podas de manejo sustentável em hidrogênio de baixo carbono.\n• Parâmetros configuráveis: fração colhível sustentável (fi) e regime tecnológico (Conservador, Típico e Otimizado).\n• Calcula a produção anual em MWh/ano e toneladas de H2."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    create_card(s7, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb2 = s7.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Rota B: Reforma a Vapor de Biogás"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_HYDRO
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "Reação de Reforma (BSR):\nCH4 + 2H2O → CO2 + 4H2 (com razão S/C)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)
    p = tf2.add_paragraph()
    p.text = "• Valorização de dejetos pecuários locais (bovinos, suínos e caprinos da Pesquisa Pecuária Municipal/IBGE).\n• Três cenários de coleta: Confinamento Técnico (15%), Semi-intensivo (40%) e Teórico Máximo (100%).\n• Guardrail Hídrico Rigoroso: Calcula a demanda exata de água (m³/ano) para garantir não concorrência com o abastecimento humano no semiárido."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 8: MÓDULO III — GOVERNANÇA PARTICIPATIVA
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Módulo III: Governança Institucional e Repartição Justa de Benefícios")
    add_footer(s8, 8)

    pillars = [
        ("8 Princípios de Elinor Ostrom", COLOR_PRIMARY, "Mapeamento da maturidade institucional comunitária (regras claras, monitoramento participativo, sanções graduadas e resolução ágil de litígios locais)."),
        ("Acordos Comunitários de Manejo", COLOR_SECONDARY, "Gestão digital de termos firmados com associações rurais, detalhando área protegida, metas de manejo e famílias beneficiadas."),
        ("Contratos de PSA e Repartição", COLOR_ACCENT, "Instrumentos de Pagamento por Serviços Ambientais com regras transparentes de rateio de receitas entre produtores, fundo comunitário e preservação.")
    ]

    left_pos = 0.8
    for title, col_color, desc in pillars:
        create_card(s8, Inches(left_pos), Inches(1.7), Inches(3.64), Inches(5.0))
        tb = s8.shapes.add_textbox(Inches(left_pos + 0.2), Inches(1.9), Inches(3.24), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(14)
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_EARTH
        p.space_after = Pt(14)
        p = tf.add_paragraph()
        p.text = "Garantia para o Financiador: Segurança jurídica e social que evita o abandono dos projetos após a fase de implantação."
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_MUTED
        left_pos += 4.04

    # ==========================================
    # SLIDE 9: MÓDULO IV — MONITORAMENTO SOCIOECOLÓGICO (MRV)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Módulo IV: Mensuração, Relato e Verificação (MRV Contínuo)")
    add_footer(s9, 9)

    create_card(s9, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb1 = s9.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Scorecard de Impacto e Metas"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    p = tf1.add_paragraph()
    p.text = "• Indicadores de Eficácia Territorial: Metas de redução de corte ilegal, área em regeneração natural assistida e volume de biomassa certificada.\n• Indicadores Sociais de Retorno: Famílias integradas a planos de manejo, geração de renda comunitária e acesso a tecnologias sociais.\n• Painel de Conformidade: Monitora desvios entre o plano operacional pactuado e a execução no terreno."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    create_card(s9, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb2 = s9.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Auditoria Participativa de Campo"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(12)
    p = tf2.add_paragraph()
    p.text = "• Diário de Bordo Digital: Registro de ocorrências, vistorias técnicas e fotos georreferenciadas diretamente pelos agentes comunitários.\n• Validação Cruzada: Combina satélites (MapBiomas Alerta) com a checagem no solo realizada pelas comunidades locais.\n• Prestação de Contas Simplificada: Dados estruturados prontos para relatórios técnicos exigidos por bancos e agências de fomento."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 10: SIMULADOR DE CENÁRIOS 10 ANOS
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Simulador de Cenários: Testando o Impacto do Investimento em 10 Anos")
    add_footer(s10, 10)

    create_card(s10, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb1 = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "4 Alavancas de Políticas Públicas"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    p = tf1.add_paragraph()
    p.text = "1. Expansão de SAFs (%): Acelera a adoção de sistemas agroflorestais sustentáveis.\n2. Redução de Lenha Predatória (%): Substituição por manejo formal e fogões ecológicos eficientes.\n3. Restauração Florestal (ha): Metas de recuperação da cobertura de Caatinga.\n4. Governança e Saneamento (%): Investimento direto em água, saneamento básico e gestão de resíduos."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    create_card(s10, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0))
    tb2 = s10.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Resultados e Curva de Trajetória"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_SECONDARY
    p.space_after = Pt(12)
    p = tf2.add_paragraph()
    p.text = "• Antes vs. Depois Instantâneo: Revela o ganho projetado no IPSE Simulado e a redução do IVS (vulnerabilidade).\n• Trajetória de 10 Anos no Gráfico: Mostra o ano exato em que o município ultrapassa a Meta Fixa de Sustentabilidade (patamar 0,60).\n• Cenários Pré-definidos: 'Transição Moderada', 'Restauração Intensiva' e 'Governança Plena' para rápida avaliação de sensibilidade."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EARTH

    # ==========================================
    # SLIDE 11: ESTUDO DE CASO MUNICIPAL
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Estudo de Caso Territorial: Os 8 Municípios-Piloto no Ceará")
    add_footer(s11, 11)

    create_card(s11, Inches(0.8), Inches(1.7), Inches(11.733), Inches(5.0))
    tb = s11.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.1), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Ranking Intermunicipal e Perfil Estratégico para Alocação de Recursos"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    municipios_info = [
        ("Tauá & Crateús (Sertões dos Inhamuns/Crateús)", "Elevadíssimo estoque de Caatinga nativa e expressivo PEL. Territórios prioritários para polos de bioenergia, biomassa certificada e conservação em larga escala."),
        ("Senador Pompeu & Quixeramobim (Sertão Central)", "Bacia leiteira expressiva e forte vulnerabilidade social (IVS). Aptos à rota integrada de biogás pecuário (Rota B de H2) com transição energética justa."),
        ("Quixadá & Sobral (Pólos Regionais)", "Maior densidade de infraestrutura agroecológica (IIE) e capacidade institucional (IGL), servindo de entrepostos logísticos e hubs tecnológicos."),
        ("Iguatu & Cascavel (Litoral/Centro-Sul)", "Perfis contrastantes de uso do solo que validam a robustez do algoritmo em cenários de transição florestal e pressão periurbana.")
    ]

    for regiao, desc in municipios_info:
        p_item = tf.add_paragraph()
        p_item.text = f"• {regiao}: {desc}"
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = COLOR_EARTH
        p_item.space_after = Pt(8)

    # ==========================================
    # SLIDE 12: CONCLUSÃO E PRÓXIMOS PASSOS
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Conclusão: Por Que Investir na Plataforma PITS?")
    add_footer(s12, 12)

    reasons = [
        ("Prontidão Tecnológica (TRL)", COLOR_PRIMARY, "A ferramenta está em operação real, com interface moderna, hot reload, arquitetura reativa e exportação instantânea de dossiês em PDF."),
        ("Escalabilidade Territorial", COLOR_SECONDARY, "Estruturada sobre dados nacionais (IBGE/MapBiomas), a metodologia pode ser facilmente expandida para todos os 184 municípios do Ceará e os 9 estados do Nordeste."),
        ("Impacto Social e Climático Mensurável", COLOR_ACCENT, "Cada investimento é monitorável via IPSE, garantindo que o recurso financeiro promova energia limpa aliada à redução da pobreza extrema.")
    ]

    left_pos = 0.8
    for title, col_color, desc in reasons:
        create_card(s12, Inches(left_pos), Inches(1.7), Inches(3.64), Inches(5.0))
        tb = s12.shapes.add_textbox(Inches(left_pos + 0.2), Inches(1.9), Inches(3.24), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(14)
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_EARTH
        left_pos += 4.04

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\Luana\.gemini\antigravity\scratch\remix-of-pixel-perfect"
    out_file = os.path.join(out_dir, "Apresentacao_PITS_Financiadores.pptx")
    create_presentation(out_file)
