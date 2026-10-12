"""
Script para download e extração dos dados municipais de cobertura do solo do MapBiomas Coleção 11.
Fonte: MapBiomas Brasil - Coleção 11 (Estatísticas Municipais)
URL oficial: https://brasil.mapbiomas.org/downloads/estatisticas/
Arquivo: MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.zip
Salva o arquivo bruto em data/raw/ e os dados filtrados em data/raw/mapbiomas_col11_municipios.json.
"""

import os
import json
import hashlib
import zipfile
import requests
import openpyxl

RAW_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
os.makedirs(RAW_DIR, exist_ok=True)

MUNICIPIOS = {
    2311306: "Quixadá",
    2304103: "Crateús",
    2305506: "Iguatu",
    2312908: "Sobral",
    2307304: "Juazeiro do Norte",
    2304202: "Crato",
    2312700: "Senador Pompeu",
    2313302: "Tauá",
    2302800: "Canindé",
    2303501: "Cascavel"
}

ZIP_URL = "https://drive.usercontent.google.com/download?id=1otOqymHuixvkRGVl65zTTNyfaHo46Gqk&export=download&confirm=t"
ZIP_PATH = os.path.join(RAW_DIR, "MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.zip")

def download_zip():
    if os.path.exists(ZIP_PATH) and os.path.getsize(ZIP_PATH) > 70000000:
        print(f"[OK] Arquivo zip já existe localmente ({os.path.getsize(ZIP_PATH)} bytes)")
        return
    print(f"-> Baixando MapBiomas Col. 11 ({ZIP_URL})...")
    r = requests.get(ZIP_URL, stream=True, headers={"User-Agent": "Mozilla/5.0"}, timeout=120)
    total_size = int(r.headers.get("content-length", 0))
    downloaded = 0
    with open(ZIP_PATH, "wb") as f:
        for chunk in r.iter_content(chunk_size=1024 * 1024): # 1MB chunks
            if chunk:
                f.write(chunk)
                downloaded += len(chunk)
                mb = downloaded / (1024 * 1024)
                print(f"   Baixado: {mb:.1f} MB...", end="\r", flush=True)
    print(f"\n[OK] Download concluído: {ZIP_PATH} ({os.path.getsize(ZIP_PATH)} bytes)")

def process_zip():
    # Hash SHA256 do arquivo zip bruto
    with open(ZIP_PATH, "rb") as f:
        sha256 = hashlib.sha256(f.read()).hexdigest()
    print(f"SHA-256 do arquivo bruto: {sha256}")

    with zipfile.ZipFile(ZIP_PATH, "r") as z:
        namelist = z.namelist()
        print(f"Arquivos no zip: {namelist}")
        excel_name = [n for n in namelist if n.endswith(".xlsx") or n.endswith(".csv")][0]
        print(f"Processando {excel_name}...")
        
        extracted_path = os.path.join(RAW_DIR, excel_name)
        if not os.path.exists(extracted_path):
            print(f"Extraindo {excel_name}...")
            z.extract(excel_name, RAW_DIR)

    # Processar o Excel
    print(f"Lendo {extracted_path} com openpyxl (read_only=True)...")
    wb = openpyxl.load_workbook(extracted_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Planilhas disponíveis: {sheet_names}")
    ws = wb['COVERAGE_11']

    # Achar colunas de geocode e anos
    header = None
    records = {code: [] for code in MUNICIPIOS.keys()}
    
    row_count = 0
    geo_col_idx = None
    for row in ws.iter_rows(values_only=True):
        row_count += 1
        if row_count == 1:
            header = [str(c).strip() if c is not None else "" for c in row]
            print(f"Colunas (total {len(header)}): {header[:15]}")
            for idx, col_name in enumerate(header):
                if any(k in col_name.lower() for k in ["geocode", "municipality_code", "cd_geocmu", "geocodigo", "codigo_ibge"]):
                    geo_col_idx = idx
                    break
            print(f"Índice coluna geocode: {geo_col_idx} ({header[geo_col_idx] if geo_col_idx is not None else 'N/A'})")
            continue
        
        geo_val = None
        if geo_col_idx is not None:
            try:
                geo_val = int(row[geo_col_idx])
            except:
                pass
        else:
            for val in row[:8]:
                try:
                    v_int = int(val)
                    if v_int in records:
                        geo_val = v_int
                        break
                except:
                    pass

        if geo_val in records:
            row_dict = {}
            for h, v in zip(header, row):
                if v is not None:
                    row_dict[h] = v
            records[geo_val].append(row_dict)

        if row_count % 50000 == 0:
            print(f"   Processadas {row_count} linhas...", flush=True)

    print(f"\n[OK] Varredura concluída. Linhas processadas: {row_count}")
    wb.close()

    output_data = {
        "fonte": "MapBiomas Brasil - Coleção 11",
        "url": "https://brasil.mapbiomas.org/downloads/estatisticas/",
        "data_consulta": "11/10/2026",
        "arquivo_bruto": os.path.basename(ZIP_PATH),
        "sha256_arquivo_bruto": sha256,
        "municipios": {}
    }

    for code, rows in records.items():
        output_data["municipios"][code] = {
            "nome": MUNICIPIOS[code],
            "total_registros_classes": len(rows),
            "dados": rows
        }
        print(f"   {MUNICIPIOS[code]} ({code}): {len(rows)} classes/registros encontrados")

    out_json = os.path.join(RAW_DIR, "mapbiomas_col11_municipios.json")
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
    print(f"[OK] Dados municipais MapBiomas salvos em {out_json}")

if __name__ == "__main__":
    download_zip()
    process_zip()
