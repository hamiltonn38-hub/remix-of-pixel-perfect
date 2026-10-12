"""
Script para download e extração dos índices do IQM (Índice de Qualidade do Meio Ambiente)
e repasses da cota-parte do ICMS Socioambiental dos 10 municípios cearenses.
Fonte oficial: IPECE / SEFAZ-CE / SEMA-CE (Governo do Estado do Ceará)
URL: https://www.ipece.ce.gov.br/cota-parte-do-icms/
"""

import os
import json
import hashlib
import requests
import openpyxl

RAW_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
os.makedirs(RAW_DIR, exist_ok=True)

MUNICIPIOS = {
    2311306: "Quixadá",
    2304103: "Crateús",
    2305506: "Iguatu",
    2312908: "Sobral",
    2307304: "Juazeiro do Norte",
    2304202: "Crato",
    2312700: "Senador Pompeu",
    2313302: "Tauá",
    2302800: "Canindé",
    2303501: "Cascavel"
}

URL_IQM_2026 = "https://www.ipece.ce.gov.br/wp-content/uploads/sites/45/2026/09/IQM_Base-de-Dados_2026.xlsx"
FILE_IQM_2026 = os.path.join(RAW_DIR, "IQM_Base-de-Dados_2026.xlsx")

URL_RESULTADOS = "https://www.ipece.ce.gov.br/wp-content/uploads/sites/45/2026/09/Resultados_2018-2026.xlsx"
FILE_RESULTADOS = os.path.join(RAW_DIR, "Resultados_2018-2026.xlsx")

def download_file(url, target_path):
    print(f"Baixando {url}...")
    r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
    with open(target_path, "wb") as f:
        f.write(r.content)
    sha256 = hashlib.sha256(r.content).hexdigest()
    print(f"[OK] Salvo em {target_path} ({len(r.content)} bytes, SHA-256: {sha256})")
    return sha256

def main():
    print("=== COLETANDO DADOS DO IQM / ICMS SOCIOAMBIENTAL (IPECE) ===")
    sha_iqm = download_file(URL_IQM_2026, FILE_IQM_2026)
    sha_res = download_file(URL_RESULTADOS, FILE_RESULTADOS)

    # Processar IQM 2026
    wb = openpyxl.load_workbook(FILE_IQM_2026, data_only=True)
    ws = wb.active

    mun_data = {code: {} for code in MUNICIPIOS.keys()}

    header_iqm = None
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 14:
            header_iqm = [str(c).strip() if c is not None else "" for c in row]
            print(f"Colunas IQM 2026: {header_iqm}")
            continue
        if row_idx > 14 and len(row) > 3 and row[3] is not None:
            mun_name = str(row[3]).strip()
            for code, name in MUNICIPIOS.items():
                if name.lower() == mun_name.lower():
                    row_dict = {
                        "municipio": mun_name,
                        "tipo_consorcio": str(row[4]).strip() if len(row) > 4 and row[4] is not None else "",
                        "eixo_I_plano": str(row[5]).strip() if len(row) > 5 and row[5] is not None else "",
                        "eixo_II_coleta": str(row[6]).strip() if len(row) > 6 and row[6] is not None else "",
                        "eixo_III_destinacao": str(row[7]).strip() if len(row) > 7 and row[7] is not None else "",
                        "eixo_IV_coleta_seletiva": str(row[8]).strip() if len(row) > 8 and row[8] is not None else "",
                        "eixo_V_logistica_reversa": str(row[9]).strip() if len(row) > 9 and row[9] is not None else "",
                        "eixo_VI_educacao_ambiental": str(row[10]).strip() if len(row) > 10 and row[10] is not None else "",
                        "iqm_somatorio": float(row[11]) if len(row) > 11 and str(row[11]).replace('.', '', 1).isdigit() else str(row[11]) if len(row) > 11 else None
                    }
                    mun_data[code]["iqm_2026"] = row_dict
                    print(f"   [IQM 2026] {name}: {row_dict['iqm_somatorio']} ({row_dict['tipo_consorcio']})")

    # Processar Resultados Históricos 2025-2026
    wb_res = openpyxl.load_workbook(FILE_RESULTADOS, data_only=True)
    ws_res = wb_res['2025-2026']
    for row_idx, row in enumerate(ws_res.iter_rows(values_only=True)):
        if row_idx >= 7 and len(row) > 2 and row[2] is not None:
            mun_name = str(row[2]).strip()
            for code, name in MUNICIPIOS.items():
                if name.lower() == mun_name.lower():
                    # Coluna 9 = IQMi, Coluna 10 = Participacao (IQMi/sum), Coluna 11 = Coeficiente Rateio (IQMi/sum * 2% * 100)
                    row_dict = {
                        "municipio": mun_name,
                        "iqm_indice": float(row[9]) if len(row) > 9 and row[9] is not None else None,
                        "participacao_meio_ambiente": float(row[10]) if len(row) > 10 and row[10] is not None else None,
                        "coeficiente_rateio_meio_ambiente_pct": float(row[11]) if len(row) > 11 and row[11] is not None else None,
                        "ano_exercicio": "2025-2026"
                    }
                    mun_data[code]["rateio_icms_socioambiental"] = row_dict
                    print(f"   [Rateio 2025-2026] {name}: Coef {row_dict['coeficiente_rateio_meio_ambiente_pct']}% (IQMi={row_dict['iqm_indice']})")


    output = {
        "fonte": "IPECE - Instituto de Pesquisa e Estratégia Econômica do Ceará",
        "url": "https://www.ipece.ce.gov.br/cota-parte-do-icms/",
        "data_consulta": "11/10/2026",
        "arquivos_brutos": {
            "IQM_Base-de-Dados_2026.xlsx": sha_iqm,
            "Resultados_2018-2026.xlsx": sha_res
        },
        "municipios": {
            str(code): {
                "nome": MUNICIPIOS[code],
                "dados": mun_data[code]
            } for code in MUNICIPIOS.keys()
        }
    }

    outpath = os.path.join(RAW_DIR, "ipece_iqm_municipios.json")
    with open(outpath, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)
    print(f"[OK] Dados IQM salvos em {outpath}")

if __name__ == "__main__":
    main()
