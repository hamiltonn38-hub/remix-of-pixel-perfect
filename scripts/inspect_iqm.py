import openpyxl
import os

RAW_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
wb = openpyxl.load_workbook(os.path.join(RAW_DIR, "IQM_Base-de-Dados_2026.xlsx"), data_only=True)
ws = wb.active
print("--- IQM 2026 Rows 10-20 ---")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if 10 <= i <= 20:
        row_str = [str(c).encode('ascii', 'replace').decode('ascii') if c is not None else "" for c in row[:15]]
        print(f"Row {i}: {row_str}")

wb2 = openpyxl.load_workbook(os.path.join(RAW_DIR, "Resultados_2018-2026.xlsx"), data_only=True)
ws2 = wb2['2025-2026']
print("--- Resultados 2025-2026 Rows 4-15 ---")
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if 4 <= i <= 15:
        row_str = [str(c).encode('ascii', 'replace').decode('ascii') if c is not None else "" for c in row[:15]]
        print(f"Row {i}: {row_str}")
